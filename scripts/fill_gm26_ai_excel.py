import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


SRC = Path("/Users/jyin/Library/CloudStorage/OneDrive-UniversityofTennessee/PES AIPSCC/2026GM-AI-related.xlsx")
RECORDS = Path("/tmp/gm26_records.json")
OUT_DIR = Path("/Users/jyin/Documents/GitHub/LLM_LTB/outputs/gm26_ai_related")
OUT = OUT_DIR / "2026GM-AI-related-filled.xlsx"
SOURCE_URL = "https://submissions.mirasmart.com/PESGM2026/Itinerary/TechnicalProgram.aspx"


AI_RE = re.compile(
    r"\bAI\b|Artificial Intelligence|Machine Learning|Generative|Neural|Foundation Models|"
    r"Reinforcement Learning|Advanced Intelligence|Explainable Artificial|AIPS|"
    r"Data\s*Center|Datacenter|Digital Twin|Predictive",
    re.I,
)


def row_text(record):
    return " ".join(str(v) for v in record.values() if v)


def autosize(ws):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 70)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_sheet(ws, headers, rows):
    ws.delete_rows(1, ws.max_row)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def main():
    records = json.loads(RECORDS.read_text())
    technical = [r for r in records if r["kind"] == "technical" and AI_RE.search(row_text(r))]
    meetings = [r for r in records if r["kind"] == "meeting" and AI_RE.search(row_text(r))]

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)
    wb = load_workbook(OUT)
    tech_ws = wb["technical sessions"]
    meeting_ws = wb["committee subcommittee WG TF "]

    write_sheet(
        tech_ws,
        ["Date", "Time", "Type", "Title", "Committee", "Room", "Source", "Filter basis"],
        [
            [
                r["date"],
                r["time"],
                r["type"],
                r["title"],
                r["committee"],
                r["room"],
                SOURCE_URL,
                "AI / ML / artificial intelligence / foundation models / data center related",
            ]
            for r in technical
        ],
    )
    write_sheet(
        meeting_ws,
        ["Track", "Date", "Time", "Title", "Room", "Source", "Filter basis"],
        [
            [
                r["track"],
                r["date"],
                r["time"],
                r["title"],
                r["room"],
                SOURCE_URL,
                "AI / ML / artificial intelligence / foundation models / data center related",
            ]
            for r in meetings
        ],
    )

    wb.save(OUT)
    print(json.dumps({"output": str(OUT), "technical_rows": len(technical), "meeting_rows": len(meetings)}, indent=2))


if __name__ == "__main__":
    main()
